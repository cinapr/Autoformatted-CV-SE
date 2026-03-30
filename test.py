from docx import Document
from copy import deepcopy
from docx.oxml.ns import qn

# ------------------------
# Load document
# ------------------------
doc = Document("EMPTY.docx")


# ------------------------
# Helper: replace text safely inside XML
# ------------------------
def replace_text_in_element(element, placeholder, value):
    for node in element.iter():
        if node.tag == qn('w:t') and node.text:
            if placeholder in node.text:
                node.text = node.text.replace(placeholder, value)

# -----------------------------------------------------------------------
# Helper: using previous font-formatting, instead of default calibri 12pt
# -----------------------------------------------------------------------
def set_paragraph_text_preserve_style(paragraph, text):
    runs = paragraph.runs

    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)



# --------------------------------------------
# Helper: replace text inline like in skills
# --------------------------------------------
def replace_inline(doc, placeholder, value):
    for para in doc.paragraphs:
        if placeholder in para.text:
            para.text = para.text.replace(placeholder, value)


# ------------------------
# Insert bullet list
# ------------------------
def insert_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder in para.text:
            parent = para._element.getparent()
            idx = parent.index(para._element)

            for item in items:
                new_para = deepcopy(para)

                # THIS is the fix
                set_paragraph_text_preserve_style(new_para, item) #new_para.text = item

                parent.insert(idx, new_para._element)
                idx += 1

            parent.remove(para._element)
            break


# ------------------------
# Insert project blocks
# ------------------------
def insert_projects(doc, projects):
    paragraphs = doc.paragraphs

    for i, para in enumerate(paragraphs):
        if "{{PROJECT_BLOCK}}" in para.text:

            parent = para._element.getparent()
            idx = parent.index(para._element)

            title_template = paragraphs[i + 1]
            bullet_template = paragraphs[i + 2]

            for proj in projects:
                # Title
                new_title = deepcopy(title_template)
                set_paragraph_text_preserve_style(new_title, proj["title"]) #new_title.text = proj["title"]
                parent.insert(idx, new_title._element)
                idx += 1

                # Bullets
                for b in proj["bullets"]:
                    new_bullet = deepcopy(bullet_template)
                    set_paragraph_text_preserve_style(new_bullet, b) #new_bullet.text = b
                    parent.insert(idx, new_bullet._element)
                    idx += 1

            # remove templates
            parent.remove(para._element)
            parent.remove(title_template._element)
            parent.remove(bullet_template._element)

            break


# ------------------------
# TEST DATA (hardcoded)
# ------------------------

wilmar_bullets = [
    "Analyzed logistics transaction data across 150+ sites using SQL",
    "Integrated reporting from 100+ factories into centralized systems",
    "Identified process bottlenecks via data analysis improving flow efficiency",
    "Reduced transaction processing time to under 3 minutes via optimization",
    "Improved SQL performance reducing timeout incidents by over 50 percent",
    "Enabled planning decisions using historical transaction trend analysis",
    "Supported logistics systems including queue loading and vehicle flow",
    "Delivered systems across 1000 plus servers in multi region operations",
    "Reduced repeated incidents by over 70 percent through root cause analysis"
]

solita_bullets = [
    "Developed REST APIs enabling structured circular waste management data exchange",
    "Translated requirements into data models and system architecture documentation",
    "Supported Agile planning with KPI driven user stories and tracking"
]

projects = [
    {
        "title": "Energy Efficiency Analysis of Collaborative Software (Python, R)",
        "bullets": [
            "Collected and analyzed energy datasets using Python and R",
            "Achieved 11.07 percent energy reduction through usage optimization insights",
            "Delivered visual actionable recommendations for users and developers"
        ]
    }
]

summary_bullets = [
    "Experienced IT consultant specializing in supply chain systems and analytics",
    "Strong background in data driven decision support and system optimization",
    "Focused on sustainable and ethical software engineering practices"
]

rnd_wilmar_bullets = [
    "Gathered transaction records from 150 factories for analysis",
    "Transformed datasets into structured decision ready visuals",
    "Built dashboards for C level stakeholders with clear insights"
]

# ------------------------
# RUN
# ------------------------

insert_bullets(doc, "{{EXP_WILMAR_ITEM}}", wilmar_bullets)
insert_bullets(doc, "{{EXP_SOLITA_ITEM}}", solita_bullets)
insert_bullets(doc, "{{SUMMARY_ITEM}}", summary_bullets)
insert_bullets(doc, "{{RND_WILMAR_ITEM}}", rnd_wilmar_bullets)

replace_inline(doc, "{{TECH_TITLE}}", "DATA")
replace_inline(doc, "{{TECH_STACK_ITEM}}", "SQL, Python, Power BI")
replace_inline(doc, "{{TECH_TITLE}}", "CODE")
replace_inline(doc, "{{TECH_STACK_ITEM}}", "C#, Python, PHP")
replace_inline(doc, "{{SKILL_TITLE}}", "ANALYSIS")
replace_inline(doc, "{{SKILL_ITEM}}", "KPI Monitoring, Process Optimization")
replace_inline(doc, "{{SKILL_TITLE}}", "SUPPORT")
replace_inline(doc, "{{SKILL_ITEM}}", "Documentation, Remote")

insert_projects(doc, projects)

# Save output
doc.save("TEST_OUTPUT.docx")

print("Done. Check TEST_OUTPUT.docx")