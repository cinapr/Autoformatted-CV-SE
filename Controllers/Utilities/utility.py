# utility.py
import sys
import json
import re
from docx import Document
from copy import deepcopy


# ------------------------
# Load file
# ------------------------
def load_file(path):
    with open(path, encoding="utf-8") as f:
        return f.read()
    


# ------------------------
# Clean Text
# ------------------------
def clean_xml_text(text):
    # Remove ANSI escape sequences (like \x1b[91m)
    text = re.sub(r'\x1b\[[0-9;]*m', '', text)

    # Remove NULL bytes and invalid XML chars
    text = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', text)
    
    # Remove other non-XML-safe control chars
    text = re.sub(r'[^\x09\x0A\x0D\x20-\x7E\u00A0-\uFFFF]', '', text)
    
    return text

def clean_dict(data):
    if isinstance(data, dict):
        return {k: clean_dict(v) for k, v in data.items()}
    elif isinstance(data, list):
        return [clean_dict(v) for v in data]
    elif isinstance(data, str):
        return clean_xml_text(data)
    return data


# ------------------------
# Save text
# ------------------------
def save_text(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)

# ------------------------
# Preserve formatting
# ------------------------
def set_text(paragraph, text):
    try:
        runs = paragraph.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
        else:
            paragraph.add_run(text)

    except Exception as e:
        print("❌ ERROR in set_text")
        print(f"Paragraph object: {paragraph}")
        print(f"Text to insert: {text}")
        print(f"Error: {str(e)}")
        raise


# ------------------------
# SIMPLE BULLETS
# ------------------------
def insert_simple_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder in para.text:
            parent = para._element.getparent()
            idx = parent.index(para._element)

            for item in items:
                new_p = deepcopy(para)
                set_text(new_p, item)
                parent.insert(idx, new_p._element)
                idx += 1

            parent.remove(para._element)
            break


# ------------------------
# LABELED BULLETS
# ------------------------
def insert_labeled_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder in para.text:

            parent = para._element.getparent()
            idx = parent.index(para._element)

            for item in items:
                new_p = deepcopy(para)

                for r in new_p.runs:
                    r.text = ""

                run_title = new_p.runs[0]
                run_title.text = item["title"].upper() + ": "

                run_value = new_p.add_run(item["value"])

                template_run = para.runs[-1]
                run_value.font.name = template_run.font.name
                run_value.font.size = template_run.font.size
                run_value.bold = False

                parent.insert(idx, new_p._element)
                idx += 1

            parent.remove(para._element)
            break



def multi_line_input_enter():
    print("Enter text (press Enter on an empty line to finish):")
    lines = []
    while True:
        line = input()
        if not line:
            break
        lines.append(line)
    text = "\n".join(lines)

    return text


def multi_line_input_INPUTFINISH():
    print("Enter your text. Type 'INPUTFINISH' on a new line to save:")

    lines = []
    while True:
        try:
            line = input()
            if line == 'INPUTFINISH':
                break
            lines.append(line)
        except EOFError:
            break

    # Join lines, excluding the INPUTFINISH line
    final_text = '\n'.join(lines)
    return final_text



def multi_line_input_INPUTFINISH_skipENTER(question_header):
    print(question_header)
    print("""
    - Press ENTER to skip
    - Or type anything to start (end with INPUTFINISH)
    """)
    
    extra_questions = ""
    lines = []
    first_line = input()

    if first_line.strip() != "":
        if first_line.strip().upper() != "INPUTFINISH":
            lines.append(first_line)

            while True:
                line = input()
                if line.strip().upper() == "INPUTFINISH":
                    break
                lines.append(line)

    extra_questions = "\n".join(lines)

    return extra_questions