import os
from docx import Document

from Controllers.Utilities.utility import load_file, save_text, multi_line_input_INPUTFINISH_skipENTER, clean_xml_text, clean_dict

def merge_documents(cv_path, cl_path, scrap_txt_path, final_output_path):
    """
    Merges CV, Cover Letter, and Scrap Text into one .docx file
    using Word sections.
    """
    # Create the master document
    master_doc = Document()

    # 1. ADD RESUME
    if os.path.exists(cv_path):
        sub_doc = Document(cv_path)
        for element in sub_doc.element.body:
            master_doc.element.body.append(element)
    
    # Add Section Break for Cover Letter
    master_doc.add_section()

    # 2. ADD COVER LETTER
    if os.path.exists(cl_path):
        sub_doc = Document(cl_path)
        for element in sub_doc.element.body:
            master_doc.element.body.append(element)

    # Add Section Break for LinkedIn Scrap
    master_doc.add_section()

    # 3. ADD LINKEDIN SCRAP (Text File)
    master_doc.add_heading('LinkedIn Job Description & Result', level=1)
    if os.path.exists(scrap_txt_path):
        with open(scrap_txt_path, 'r', encoding='utf-8') as f:
            for line in f:
                try:
                    master_doc.add_paragraph(clean_xml_text(line.strip()))
                except Exception as e:
                    print("Failed line: " + line.strip())
                    print(f"Exception: {e}")

    # Save the final merged document
    master_doc.save(final_output_path)
    print(f"✅ Merged document created: {final_output_path}")