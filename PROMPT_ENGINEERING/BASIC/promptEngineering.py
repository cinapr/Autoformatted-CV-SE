# pip install python-docx google-generativeai

'''
To maintain a true context window across all three phases, we need to use Gemini's start_chat() feature. 
This preserves the exact history of the conversation, so Phase 2 and Phase 3 can naturally refer back to the job details or any analysis done in Phase 1.

Important points:
1. model.start_chat(): This initialises the state tracker.

2. chat.send_message(): Replaced the standalone generation calls. Every time this function runs, it appends both your prompt and the model's response to an underlying array, giving the model access to everything that happened in previous steps.

3. Cleaner Prompts: Notice how the Phase 2 prompt no longer needs you to manually extract and inject the keywords via Python strings ({keywords}). The model simply looks back at its own previous response to find them.
'''

import os
import json
import google.generativeai as genai
from docx import Document

# Setup your API key from the environment variable
api_key = os.environ.get("GEMINI_API_KEY")
if not api_key:
    raise ValueError("GEMINI_API_KEY environment variable not set.")

genai.configure(api_key=api_key)

MASTER_PROFILE = """
John Doe
Email: john.doe@email.com | Phone: 555-0100
Software Engineer with 5 years of experience building scalable backend systems.
Skills: Python, React, PostgreSQL, Docker, AWS, CI/CD.
Experience:
- Backend Engineer at TechCorp (2021-Present): Reduced API latency by 40% using Redis. Built microservices in Python.
- Junior Developer at WebSolutions (2019-2021): Managed database migrations and wrote unit tests.
"""

JOB_DESCRIPTION = """
We are looking for a Senior Python Developer to join our core infrastructure team.
Requirements:
- 4+ years of Python experience.
- Deep knowledge of AWS services (EC2, S3, RDS).
- Experience transitioning monolithic architectures to microservices.
- Must be located in North America or willing to relocate.
"""

def create_word_document(resume_data, filename="Tailored_Resume.docx"):
    doc = Document()
    doc.add_heading(resume_data.get("name", "Name Missing"), 0)
    doc.add_paragraph(resume_data.get("contact_info", ""))
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(resume_data.get("summary", ""))
    
    doc.add_heading("Experience", level=1)
    for job in resume_data.get("experience", []):
        job_heading = f"{job.get('title', '')} at {job.get('company', '')}"
        doc.add_heading(job_heading, level=2)
        for bullet in job.get("bullets", []):
            doc.add_paragraph(bullet, style='List Bullet')
            
    doc.add_heading("Technical Skills", level=1)
    skills_text = ", ".join(resume_data.get("skills", []))
    doc.add_paragraph(skills_text)
    
    doc.save(filename)
    print(f"Saved resume to {filename}")

def run_job_pipeline():
    system_instruction = f"""
    You are an expert career strategist and technical resume writer. 
    You process job descriptions and tailor applicant profiles to match. 
    Do not invent experience. Rely strictly on the provided Master Profile.
    
    Master Profile:
    {MASTER_PROFILE}
    """

    model = genai.GenerativeModel(
        model_name="gemini-1.5-pro",
        system_instruction=system_instruction
    )

    # Start a chat session to maintain the context window automatically
    chat = model.start_chat(history=[])

    print("Phase 1: Analyzing job suitability...")
    
    suitability_prompt = f"""
    Evaluate this job description against the Master Profile. 
    Return a strict JSON object with these keys: 
    - "match_score" (integer 0-100)
    - "apply_recommendation" (boolean)
    - "missing_skills" (list of strings)
    - "visa_location_blockers" (string)
    - "targeted_keywords" (list of strings)
    
    Job Description:
    {JOB_DESCRIPTION}
    """
    
    # Use chat.send_message instead of model.generate_content
    response_1 = chat.send_message(
        suitability_prompt,
        generation_config={"response_mime_type": "application/json"}
    )
    
    suitability_data = json.loads(response_1.text)
    print(f"Match Score: {suitability_data.get('match_score')}%")
    
    if not suitability_data.get("apply_recommendation"):
        print("\nPipeline stopped. The AI does not recommend applying for this role.")
        return

    print("\nPhase 2: Generating tailored resume data...")
    
    resume_prompt = """
    Based on your previous evaluation, rewrite the Master Profile to target the job description. 
    Focus heavily on the targeted keywords you identified. 
    Return a strict JSON object matching this exact schema:
    {
        "name": "string",
        "contact_info": "string",
        "summary": "string",
        "experience": [
            {
                "company": "string",
                "title": "string",
                "bullets": ["string", "string"]
            }
        ],
        "skills": ["string"]
    }
    Keep all bullets concise and action-oriented.
    """
    
    response_2 = chat.send_message(
        resume_prompt,
        generation_config={"response_mime_type": "application/json"}
    )
    
    resume_data = json.loads(response_2.text)
    create_word_document(resume_data)
    
    print("\nPhase 3: Drafting the cover letter...")
    
    cover_letter_prompt = """
    Write a 3-paragraph cover letter for this role. 
    Address the specific goals of the company mentioned in the job description you analyzed in the first turn. 
    Highlight past achievements from the Master Profile that solve those specific problems. 
    Do not use generic opening hooks. Output plain text.
    """
    
    # We drop the JSON restriction here to get natural paragraphs
    response_3 = chat.send_message(cover_letter_prompt)
    
    with open("Cover_Letter.txt", "w") as f:
        f.write(response_3.text)
    print("Saved cover letter to Cover_Letter.txt")

if __name__ == "__main__":
    run_job_pipeline()