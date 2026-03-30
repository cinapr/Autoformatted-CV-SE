from docx import Document
from copy import deepcopy
import json

doc = Document("TEMPLATE.docx")


# ------------------------
# Preserve formatting
# ------------------------
def set_text(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


# ------------------------
# BULLET SECTION (generic)
# ------------------------
def insert_simple_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder in para.text:
            parent = para._element.getparent()
            idx = parent.index(para._element)

            for item in items:
                new_p = deepcopy(para)
                set_text(new_p, item)
                parent.insert(idx, new_p._element)
                idx += 1

            parent.remove(para._element)
            break


# --------------------------------
# LINES BULLET SECTION (title + value)
# --------------------------------
def insert_labeled_bullets(doc, placeholder, items):
    for para in doc.paragraphs:
        if placeholder in para.text:

            parent = para._element.getparent()
            idx = parent.index(para._element)

            for item in items:
                new_p = deepcopy(para)

                # clear runs
                for r in new_p.runs:
                    r.text = ""

                # title (bold)
                run_title = new_p.runs[0]
                run_title.text = item["title"].upper() + ": "

                # value (normal)
                #new_p.add_run(item["value"])
                run_value = new_p.add_run(item["value"])

                # copy font style from template
                template_run = para.runs[-1]
                run_value.font.name = template_run.font.name
                run_value.font.size = template_run.font.size
                run_value.bold = False

                parent.insert(idx, new_p._element)
                idx += 1

            parent.remove(para._element)
            break


# ------------------------
# EXPERIENCE BLOCK
# ------------------------
def insert_experience(doc, experiences):
    paras = doc.paragraphs

    for i, para in enumerate(paras):
        if "{{EXPERIENCE_BLOCK}}" in para.text:

            parent = para._element.getparent()
            idx = parent.index(para._element)

            title_template = paras[i + 1]
            bullet_template = paras[i + 2]
            project_title_template = paras[i + 3]
            project_bullet_template = paras[i + 4]

            for exp in experiences:
                # Title line
                title_text = f"{exp['title']} | {exp['company']} | {exp['location']} | {exp['time']}"
                title_p = deepcopy(title_template)
                set_text(title_p, title_text)
                parent.insert(idx, title_p._element)
                idx += 1

                # Main bullets
                for b in exp["bullets"]:
                    bullet_p = deepcopy(bullet_template)
                    set_text(bullet_p, b)
                    parent.insert(idx, bullet_p._element)
                    idx += 1

                # Optional project section
                if exp.get("projects"):
                    proj_title = deepcopy(project_title_template)
                    parent.insert(idx, proj_title._element)
                    idx += 1

                    for pb in exp["projects"]:
                        proj_b = deepcopy(project_bullet_template)
                        set_text(proj_b, pb)
                        parent.insert(idx, proj_b._element)
                        idx += 1

                # AFTER finishing one experience give newline spaces in between
                spacer = deepcopy(title_template)
                set_text(spacer, "")
                parent.insert(idx, spacer._element)
                idx += 1

            # cleanup
            parent.remove(para._element)
            parent.remove(title_template._element)
            parent.remove(bullet_template._element)
            parent.remove(project_title_template._element)
            parent.remove(project_bullet_template._element)

            break


# ------------------------
# PROJECT BLOCK
# ------------------------
def insert_projects(doc, projects):
    paras = doc.paragraphs

    for i, para in enumerate(paras):
        if "{{PROJECT_BLOCK}}" in para.text:

            parent = para._element.getparent()
            idx = parent.index(para._element)

            title_template = paras[i + 1]
            bullet_template = paras[i + 2]

            for proj in projects:
                title_p = deepcopy(title_template)
                set_text(title_p, proj["title"])
                parent.insert(idx, title_p._element)
                idx += 1

                for b in proj["bullets"]:
                    bullet_p = deepcopy(bullet_template)
                    set_text(bullet_p, b)
                    parent.insert(idx, bullet_p._element)
                    idx += 1

                # AFTER finishing one project give newline spaces in between
                spacer = deepcopy(title_template)
                set_text(spacer, "")
                parent.insert(idx, spacer._element)
                idx += 1

            parent.remove(para._element)
            parent.remove(title_template._element)
            parent.remove(bullet_template._element)

            break

# ----------------------------------------
# REPLACE SINGLE VALUE (FOCUS/THESIS TITLE)
# -----------------------------------------
def replace_single_value(doc, placeholder, value):
    for para in doc.paragraphs:
        if placeholder in para.text:
            for run in para.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

# ------------------------
# REMOVE OPTIONAL SECTION
# ------------------------
def remove_section(doc, section_title):
    remove = False
    for para in doc.paragraphs:
        if section_title in para.text:
            remove = True

        if remove:
            p = para._element
            p.getparent().remove(p)

        if remove and "EDUCATION" in para.text:
            break


# ------------------------
# LOAD JSON (example)
# ------------------------

with open("data.json", encoding="utf-8") as f:
    data = json.load(f)

# ------------------------
# RUN
# ------------------------

insert_simple_bullets(doc, "{{SUMMARY_ITEM}}", data["summary"])

insert_labeled_bullets(doc, "{{TECH_TITLE}}", data["tech"])
insert_labeled_bullets(doc, "{{SKILL_TITLE}}", data["skills"])

insert_simple_bullets(doc, "{{AWARDS_BULLET}}", data["awards"])

insert_experience(doc, data["experience"])

replace_single_value(doc, "{{FOCUS_REPLACE}}", data["education"]["focus"])
replace_single_value(doc, "{{THESIS_REPLACE}}", data["education"]["thesis"])

if data.get("projects"):
    insert_projects(doc, data["projects"])
else:
    remove_section(doc, "OPEN-SOURCE PROJECTS")

doc.save("FINAL_OUTPUT.docx")